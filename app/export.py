"""Ersetzung anwenden und Ergebnis als TXT, DOCX oder PDF ausgeben.

Alle Exporte werden ohne Metadaten geschrieben. Ein anonymisiertes Dokument,
in dessen Autorfeld noch der Klarname steht, ist nicht anonymisiert.
"""

from __future__ import annotations

import contextlib
import html
import io
from datetime import UTC, datetime

from docx import Document
from docx.shared import Pt
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate


def ersetze(text: str, ersetzungen: list[dict]) -> str:
    """Wendet Span-Ersetzungen an. Rein mechanisch, keine eigene Logik.

    Von hinten nach vorn, damit frühere Offsets gültig bleiben.
    """
    gueltig = [e for e in ersetzungen
               if 0 <= e["start"] < e["ende"] <= len(text)]
    for eintrag in sorted(gueltig, key=lambda e: e["start"], reverse=True):
        text = text[:eintrag["start"]] + eintrag["label"] + text[eintrag["ende"]:]
    return text


# Die Stufe gehört in den Vermerk: wer das Dokument später in die Hand nimmt,
# muss erkennen können, ob er ein anonymisiertes oder ein pseudonymisiertes
# Dokument vor sich hat. "Leicht" bleibt personenbezogen.
#
# Der Haftungshinweis stand bisher nur im Bestätigungsdialog vor dem Export
# (static/app.js) — den sieht niemand mehr, sobald die Datei die Anwendung
# verlassen hat. Wer sie Monate später weiterreicht, ausdruckt oder
# archiviert, hat nur noch das, was im Dokument selbst steht. Deshalb gehört
# der Satz hierher, nicht nur in den Dialog.
_VERMERK = {
    "leicht": "Pseudonymisiert (Stufe leicht) am {datum} — enthält weiterhin "
              "Personenbezug, vertraulich behandeln. Vollständigkeit der "
              "Pseudonymisierung ist nicht garantiert.",
    "standard": "Anonymisiert (Stufe Standard) am {datum} — automatisch "
                "erkannt und manuell geprüft. Vollständigkeit ist nicht "
                "garantiert, die Verantwortung für die Weitergabe liegt "
                "beim Absender.",
    "streng": "Anonymisiert (Stufe streng, ohne Zeitbezug) am {datum} — "
              "automatisch erkannt und manuell geprüft. Vollständigkeit ist "
              "nicht garantiert, die Verantwortung für die Weitergabe liegt "
              "beim Absender.",
}


def _kopfzeile(stufe: str = "standard") -> str:
    vorlage = _VERMERK.get(stufe, _VERMERK["standard"])
    return vorlage.format(datum=f"{datetime.now(tz=UTC).astimezone():%d.%m.%Y}")


def als_txt(text: str, mit_kopfzeile: bool = True, stufe: str = "standard") -> bytes:
    inhalt = (f"{_kopfzeile(stufe)}\n{'-' * 70}\n\n{text}\n"
              if mit_kopfzeile else text)
    return inhalt.encode("utf-8")


def als_docx(text: str, mit_kopfzeile: bool = True, stufe: str = "standard") -> bytes:
    dokument = Document()

    stil = dokument.styles["Normal"]
    stil.font.name = "Calibri"
    stil.font.size = Pt(11)
    stil.paragraph_format.space_after = Pt(8)
    stil.paragraph_format.line_spacing = 1.15

    if mit_kopfzeile:
        kopf = dokument.add_paragraph(_kopfzeile(stufe))
        kopf.runs[0].font.size = Pt(8)
        kopf.runs[0].font.italic = True

    for block in text.split("\n\n"):
        if not block.strip():
            continue
        absatz = dokument.add_paragraph()
        # Zeilenumbrüche innerhalb eines Absatzes sind nach dem Reflow
        # gewollt — Anschriftenblock, Unterschriftenzeile, Aufzählung.
        zeilen = block.split("\n")
        for i, zeile in enumerate(zeilen):
            lauf = absatz.add_run(zeile)
            if i < len(zeilen) - 1:
                lauf.add_break()

    kern = dokument.core_properties
    for feld in ("author", "last_modified_by", "title", "subject",
                 "comments", "category", "keywords", "identifier"):
        # Nicht jedes Feld ist in jeder python-docx-Fassung beschreibbar.
        with contextlib.suppress(AttributeError, ValueError):
            setattr(kern, feld, "")

    puffer = io.BytesIO()
    dokument.save(puffer)
    return puffer.getvalue()


def als_pdf(text: str, mit_kopfzeile: bool = True, stufe: str = "standard") -> bytes:
    puffer = io.BytesIO()
    doc = SimpleDocTemplate(
        puffer, pagesize=A4,
        leftMargin=25 * mm, rightMargin=20 * mm,
        topMargin=20 * mm, bottomMargin=20 * mm,
        title="", author="", subject="", creator="",
    )

    koerper = ParagraphStyle(
        "Koerper", fontName="Helvetica", fontSize=10.5, leading=15.5,
        alignment=TA_LEFT, spaceAfter=9,
    )
    notiz = ParagraphStyle(
        "Notiz", fontName="Helvetica-Oblique", fontSize=7.5,
        leading=11, textColor="#666666", spaceAfter=10,
    )

    fluss = []
    if mit_kopfzeile:
        fluss.append(Paragraph(html.escape(_kopfzeile(stufe)), notiz))

    for block in text.split("\n\n"):
        if not block.strip():
            continue
        # Der Absatz fliesst über die volle Satzbreite; nur die verbliebenen
        # Umbrüche innerhalb des Blocks werden erzwungen.
        fluss.append(Paragraph(html.escape(block).replace("\n", "<br/>"), koerper))

    doc.build(fluss)
    return puffer.getvalue()


FORMATE = {
    "txt": (als_txt, "text/plain; charset=utf-8", "txt"),
    "docx": (als_docx,
             "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
             "docx"),
    "pdf": (als_pdf, "application/pdf", "pdf"),
}

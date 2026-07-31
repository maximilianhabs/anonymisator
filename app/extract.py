"""Textextraktion aus PDF, DOCX und Klartext.

Alles läuft lokal. Kein Aufruf verlässt den Rechner.
"""

from __future__ import annotations

import io
import re
from dataclasses import dataclass, field

import fitz  # PyMuPDF
import pytesseract
from PIL import Image

# Unter diesem Schnitt an Zeichen pro Seite gehen wir davon aus, dass das PDF
# keinen brauchbaren Textlayer hat und gescannt wurde.
MIN_ZEICHEN_PRO_SEITE = 100

# 300 dpi ist der übliche Kompromiss: darunter leidet die OCR-Qualität bei
# kleinen Serifenschriften spürbar, darüber wird es nur noch langsamer.
OCR_DPI = 300


@dataclass
class Extraktion:
    text: str
    quelle: str  # "pdf-textlayer" | "pdf-ocr" | "docx" | "paste"
    seiten: int = 0
    hinweise: list[str] = field(default_factory=list)


# Zeilenanfänge, die immer einen neuen Absatz beginnen und deshalb nie an die
# vorherige Zeile angehängt werden dürfen.
_AUFZAEHLUNG = re.compile(r"^\s*(?:[-–—•*·>]|\(?\d{1,2}[.)]|[a-z][.)])\s+")


def _volle_breite(zeilen: list[str]) -> float:
    """Typische Satzbreite des Dokuments, robust gegen einzelne Ausreisser."""
    laengen = sorted(len(z) for z in zeilen if z.strip())
    if len(laengen) < 3:
        return 0.0
    return laengen[int(len(laengen) * 0.9)] if len(laengen) > 9 else laengen[-1]


def _reflow(text: str, mindestanteil: float = 0.25) -> str:
    """Fügt weich umbrochene Zeilen wieder zu durchgehenden Absätzen zusammen.

    Ein PDF kennt keine Absätze, nur Zeilen. Wer diese Zeilenenden als feste
    Umbrüche übernimmt, bekommt sie beim Export erneut aufgezwungen — bei
    anderer Satzbreite entstehen daraus Flatterzeilen und einzelne Wörter auf
    eigener Zeile. Beim Anonymisieren verschärft sich das noch, weil ein
    Pseudonym selten genauso lang ist wie der ersetzte Klarname.

    Erkennungsmerkmal ist die Zeilenlänge: eine Zeile, die bis an den Rand
    reicht, wurde vom Satz umbrochen und gehört mit der nächsten zusammen.
    Eine kürzere Zeile beendet den Absatz. Anschriftenblöcke, Unterschriften
    und Aufzählungen bestehen aus kurzen Zeilen und bleiben damit von allein
    unangetastet.
    """
    zeilen = text.split("\n")
    breite = _volle_breite(zeilen)
    if breite < 45:
        return text

    schwelle = breite * 0.8
    inhalt = [z for z in zeilen if z.strip()]
    if not inhalt:
        return text

    anteil = sum(1 for z in inhalt if len(z) >= schwelle) / len(inhalt)
    if anteil < mindestanteil:
        # Sieht nicht nach hartem Umbruch aus, sondern nach Text, in dem jede
        # Zeile bereits ein eigener Absatz ist. Dann nichts anfassen.
        return text

    ergebnis: list[str] = []
    vorherige_voll = False
    for zeile in zeilen:
        gefuellt = bool(zeile.strip()) and len(zeile) >= schwelle
        if (ergebnis and vorherige_voll and zeile.strip()
                and not _AUFZAEHLUNG.match(zeile)):
            ergebnis[-1] = f"{ergebnis[-1].rstrip()} {zeile.lstrip()}"
        else:
            ergebnis.append(zeile)
        vorherige_voll = gefuellt

    return "\n".join(ergebnis)


def _normalisiere(text: str) -> str:
    """Vereinheitlicht Zeilenenden und Sonderzeichen.

    Muss vor der Erkennung passieren, damit alle Offsets auf denselben
    Textstand zeigen.
    """
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    # Weiche Trennstriche und geschützte Leerzeichen aus PDFs
    text = text.replace("­", "").replace(" ", " ")
    # Typografische Anführungszeichen und Bindestriche vereinheitlichen,
    # sonst greifen die Regex-Regeln an manchen Stellen nicht.
    for alt, neu in [("‘", "'"), ("’", "'"), ("“", '"'),
                     ("”", '"'), ("„", '"'), ("–", "-"),
                     ("—", "-"), ("−", "-")]:
        text = text.replace(alt, neu)
    # Trennung am Zeilenende zusammenziehen ("Kopf-\nschmerz" -> "Kopfschmerz")
    text = re.sub(r"(\w)-\n(\w)", r"\1\2", text)
    # Mehr als zwei Leerzeilen bringen nichts
    text = re.sub(r"\n{3,}", "\n\n", text)
    text = re.sub(r"[ \t]{2,}", " ", text)
    return text.strip()


def _pdf_absaetze(doc) -> str:
    """Liest den Text absatzweise statt zeilenweise.

    get_text("blocks") liefert die Layoutblöcke des PDFs; bei Fliesstext
    entspricht ein Block einem Absatz. Der einfache Textmodus verliert diese
    Grenze vollständig, weil Absatzabstände im PDF blosser Weissraum sind und
    keinen Text erzeugen — der ganze Brief käme als ein Klumpen an, in dem
    sich Anrede, Anamnese und Befund nicht mehr trennen lassen.
    """
    absaetze: list[str] = []
    for seite in doc:
        for block in seite.get_text("blocks", sort=True):
            # block: (x0, y0, x1, y1, text, nummer, typ) — Typ 1 sind Bilder
            if len(block) >= 7 and block[6] != 0:
                continue
            inhalt = block[4].strip()
            if inhalt:
                absaetze.append(inhalt)
    return "\n\n".join(absaetze)


def aus_pdf(daten: bytes) -> Extraktion:
    hinweise: list[str] = []
    with fitz.open(stream=daten, filetype="pdf") as doc:
        seiten = doc.page_count
        text = _pdf_absaetze(doc)

        zeichen_pro_seite = len(text.strip()) / max(seiten, 1)
        if zeichen_pro_seite >= MIN_ZEICHEN_PRO_SEITE:
            meta_treffer = _pdf_metadaten_pruefen(doc)
            if meta_treffer:
                hinweise.append(
                    "PDF-Metadaten enthalten Angaben (" + ", ".join(meta_treffer)
                    + "). Sie werden beim Export nicht übernommen."
                )
            return Extraktion(_reflow(_normalisiere(text)),
                              "pdf-textlayer", seiten, hinweise)

        # Kein oder zu wenig Textlayer -> gescanntes Dokument, OCR nötig.
        hinweise.append(
            f"Kein verwertbarer Textlayer ({zeichen_pro_seite:.0f} Zeichen/Seite). "
            f"{seiten} Seite(n) wurden per OCR gelesen — Lesefehler sind möglich, "
            "bitte den Text im Review gegenlesen."
        )
        ocr_teile = []
        for seite in doc:
            pix = seite.get_pixmap(dpi=OCR_DPI)
            bild = Image.open(io.BytesIO(pix.tobytes("png")))
            ocr_teile.append(pytesseract.image_to_string(bild, lang="deu"))
        return Extraktion(_reflow(_normalisiere("\n\n".join(ocr_teile))),
                          "pdf-ocr", seiten, hinweise)


def _pdf_metadaten_pruefen(doc) -> list[str]:
    """Meldet befüllte Metadatenfelder, die Klarnamen enthalten können."""
    treffer = []
    for feld in ("author", "title", "subject", "creator", "producer", "keywords"):
        wert = (doc.metadata or {}).get(feld) or ""
        if wert.strip() and feld in ("author", "title", "subject", "keywords"):
            treffer.append(f"{feld}: {wert.strip()[:40]}")
    return treffer


def aus_docx(daten: bytes) -> Extraktion:
    import docx  # python-docx

    dokument = docx.Document(io.BytesIO(daten))
    teile: list[str] = []

    # Kopf- und Fußzeilen zuerst: dort stehen bei Arztbriefen fast immer
    # Klinikname, Anschrift und Telefonnummern. Wer nur document.paragraphs
    # liest, übersieht genau diesen Block.
    for abschnitt in dokument.sections:
        for bereich in (abschnitt.header, abschnitt.footer):
            for absatz in bereich.paragraphs:
                if absatz.text.strip():
                    teile.append(absatz.text)

    for absatz in dokument.paragraphs:
        teile.append(absatz.text)

    for tabelle in dokument.tables:
        for zeile in tabelle.rows:
            zellen = [z.text.strip() for z in zeile.cells]
            # Verbundene Zellen liefert python-docx mehrfach zurück
            entdoppelt = [z for i, z in enumerate(zellen) if i == 0 or z != zellen[i - 1]]
            if any(entdoppelt):
                teile.append("\t".join(entdoppelt))

    hinweise = []
    kern = dokument.core_properties
    meta = [f for f in ("author", "last_modified_by", "title", "subject")
            if (getattr(kern, f, None) or "")]
    if meta:
        hinweise.append(
            "DOCX-Metadaten enthalten Angaben (" + ", ".join(meta)
            + "). Sie werden beim Export nicht übernommen."
        )

    return Extraktion(_normalisiere("\n".join(teile)), "docx", 0, hinweise)


def aus_upload(dateiname: str, daten: bytes) -> Extraktion:
    name = dateiname.lower()
    if name.endswith(".pdf"):
        return aus_pdf(daten)
    if name.endswith(".docx"):
        return aus_docx(daten)
    if name.endswith(".doc"):
        raise ValueError(
            "Das alte .doc-Format wird nicht unterstützt. "
            "Bitte in Word als .docx speichern oder den Text direkt einfügen."
        )
    if name.endswith(".txt"):
        return Extraktion(_normalisiere(daten.decode("utf-8", errors="replace")), "paste")
    raise ValueError(f"Dateityp nicht unterstützt: {dateiname}")


def aus_text(text: str) -> Extraktion:
    # Eingefügter Text stammt oft aus einem PDF-Betrachter und trägt dessen
    # harte Zeilenumbrüche mit. _reflow greift nur, wenn der Text tatsächlich
    # so aussieht, und lässt sauber gesetzten Text unangetastet.
    return Extraktion(_reflow(_normalisiere(text)), "paste")

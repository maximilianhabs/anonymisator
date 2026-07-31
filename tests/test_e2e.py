"""End-to-End über die laufende API: PDF, DOCX, OCR, alle Exportformate.

Der harte Test steht am Ende: kommt in der fertigen Exportdatei — Binärinhalt
und Metadaten eingeschlossen — noch ein Klarname vor?
"""
import io
import re
import sys
import zipfile

import fitz
import requests
from docx import Document
from reportlab.lib.pagesizes import A4
from reportlab.lib.units import mm
from reportlab.pdfgen import canvas

BASIS = "http://127.0.0.1:5302"

BRIEF = """Neurologische Klinik am Auwald
Ludwigstrasse 47, 94032 Passau
Tel.: 0851 / 4409-233
E-Mail: sekretariat@klinik-auwald.de

Passau, 14.03.2026

Sehr geehrte Frau Kollegin Brandner,

wir berichten ueber Ihre Patientin Frau Elisabeth Hartmann, geb. 22.07.1958,
wohnhaft Rosenweg 12, 94469 Deggendorf, Versichertennummer M123456789.
Fallnummer: 2026-KA-88213

Die Vorbefunde von Dr. Sailer aus Vilshofen lagen vor.
Das MRT zeigte eine Laesion im linken Temporallappen.

Mit freundlichen Gruessen
Prof. Dr. med. Konrad Weissbach"""

# Namen, die nach der Anonymisierung nirgends mehr auftauchen duerfen.
KLARNAMEN = ["Hartmann", "Weissbach", "Brandner", "Sailer", "M123456789",
             "22.07.1958", "Rosenweg", "sekretariat@klinik-auwald.de",
             "2026-KA-88213", "0851"]

fehler = []


def pruefe(bedingung, beschreibung):
    print(f"  {'OK  ' if bedingung else 'FEHL'}  {beschreibung}")
    if not bedingung:
        fehler.append(beschreibung)


def baue_pdf(text, autor="Dr. Konrad Weissbach"):
    puffer = io.BytesIO()
    c = canvas.Canvas(puffer, pagesize=A4)
    c.setAuthor(autor)
    c.setTitle("Arztbrief Hartmann, Elisabeth")
    y = 270 * mm
    for zeile in text.split("\n"):
        c.setFont("Helvetica", 10)
        c.drawString(20 * mm, y, zeile)
        y -= 5.5 * mm
    c.save()
    return puffer.getvalue()


def baue_scan_pdf(text):
    """PDF ganz ohne Textlayer — nur gerasterte Seitenbilder."""
    quelle = fitz.open(stream=baue_pdf(text), filetype="pdf")
    ziel = fitz.open()
    for seite in quelle:
        pix = seite.get_pixmap(dpi=200)
        neu = ziel.new_page(width=seite.rect.width, height=seite.rect.height)
        neu.insert_image(neu.rect, stream=pix.tobytes("png"))
    return ziel.tobytes()


def baue_docx(text):
    dok = Document()
    dok.core_properties.author = "Konrad Weissbach"
    abschnitt = dok.sections[0]
    abschnitt.header.paragraphs[0].text = "Neurologische Klinik am Auwald - Ludwigstrasse 47, 94032 Passau"
    for absatz in text.split("\n"):
        dok.add_paragraph(absatz)
    puffer = io.BytesIO()
    dok.save(puffer)
    return puffer.getvalue()


def analysiere_datei(name, daten):
    antwort = requests.post(f"{BASIS}/api/analysieren/datei",
                            files={"datei": (name, daten)}, timeout=300)
    antwort.raise_for_status()
    return antwort.json()


def text_aus(daten, format_):
    """Holt den lesbaren Text aus einer Exportdatei zurueck."""
    if format_ == "txt":
        return daten.decode("utf-8")
    if format_ == "pdf":
        with fitz.open(stream=daten, filetype="pdf") as doc:
            meta = " ".join(str(v) for v in (doc.metadata or {}).values())
            return meta + " " + "\n".join(s.get_text() for s in doc)
    if format_ == "docx":
        # Alle XML-Teile lesen, damit auch Metadaten erfasst werden.
        with zipfile.ZipFile(io.BytesIO(daten)) as z:
            return " ".join(z.read(n).decode("utf-8", "replace")
                            for n in z.namelist() if n.endswith(".xml"))
    raise ValueError(format_)


print("\n=== 1. PDF mit Textlayer ===")
daten = analysiere_datei("brief.pdf", baue_pdf(BRIEF))
pruefe(daten["quelle"] == "pdf-textlayer", f"Quelle erkannt: {daten['quelle']}")
pruefe(len(daten["gruppen"]) > 12, f"{len(daten['gruppen'])} Gruppen gefunden")
pruefe(any("Metadaten" in h for h in daten["hinweise"]),
       "Warnt vor Klarnamen in den PDF-Metadaten")

print("\n=== 2. DOCX inkl. Kopfzeile ===")
docx_daten = analysiere_datei("brief.docx", baue_docx(BRIEF))
pruefe(docx_daten["quelle"] == "docx", f"Quelle erkannt: {docx_daten['quelle']}")
pruefe("Ludwigstrasse 47" in docx_daten["text"],
       "Kopfzeile wurde mitgelesen (dort steht die Klinikanschrift)")

print("\n=== 3. Gescanntes PDF (nur Bild, OCR noetig) ===")
ocr = analysiere_datei("scan.pdf", baue_scan_pdf(BRIEF))
pruefe(ocr["quelle"] == "pdf-ocr", f"Quelle erkannt: {ocr['quelle']}")
pruefe("Hartmann" in ocr["text"], "OCR hat den Patientennamen gelesen")
erkannt = {g["anzeige"] for g in ocr["gruppen"]}
pruefe(any("Hartmann" in e for e in erkannt),
       f"Patientenname auch im OCR-Text erkannt ({len(ocr['gruppen'])} Gruppen)")

print("\n=== 4. Exporte — bleibt irgendwo ein Klarname stehen? ===")
ersetzungen = [{"start": f["start"], "ende": f["ende"], "label": g["label"]}
               for g in daten["gruppen"] for f in g["funde"]]

for format_ in ("txt", "docx", "pdf"):
    antwort = requests.post(f"{BASIS}/api/export", json={
        "text": daten["text"], "ersetzungen": ersetzungen,
        "format": format_, "dateiname": "brief", "mit_kopfzeile": True,
    }, timeout=120)
    antwort.raise_for_status()
    inhalt = antwort.content
    lesbar = text_aus(inhalt, format_)

    print(f"\n  --- {format_.upper()} ({len(inhalt):,} Bytes) ---")
    pruefe(len(inhalt) > 200, "Datei ist nicht leer")
    pruefe(bool(re.search(r"\[[A-Z\u00c4\u00d6\u00dc]+_[A-Z0-9]+\]", lesbar)), "Pseudonyme sind enthalten")
    pruefe("[PATIENT_A]" in lesbar, "Patientin ist als PATIENT eingestuft")
    uebrig = [n for n in KLARNAMEN if n in lesbar]
    pruefe(not uebrig, f"Kein Klarname mehr enthalten{' — GEFUNDEN: ' + str(uebrig) if uebrig else ''}")

print("\n" + "=" * 60)
if fehler:
    print(f"{len(fehler)} FEHLER:")
    for f in fehler:
        print("  -", f)
    sys.exit(1)
print("Alle Pruefungen bestanden.")
